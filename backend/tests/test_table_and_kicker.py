"""Tests for elevated table styling, horizontal divider rules, and kicker accent dot treatments."""
from pptx import Presentation
from htmlppt.analyze import ir
from htmlppt.brand.presets import apply_preset, palette_brand
from htmlppt.layout.engine import layout_slide, _build_table_element
from htmlppt.layout.grid import SlideGeometry
from htmlppt.layout import model as lm
from htmlppt.render.pptx_writer import write_pptx
from htmlppt.render.tables import render_table_element


def test_table_styling_modern_rules_and_headers():
    brand = palette_brand("ocean")
    brand = apply_preset(brand, "corporate")
    geom = SlideGeometry.from_brand(brand, has_title=True)
    cr = geom.content_rect

    # 3x3 table with header row
    headers = [ir.Cell(runs=[ir.Run(t=f"Col {ci}")], header=True) for ci in range(3)]
    row1 = [ir.Cell(runs=[ir.Run(t=f"Val 1-{ci}")]) for ci in range(3)]
    row2 = [ir.Cell(runs=[ir.Run(t=f"Val 2-{ci}")]) for ci in range(3)]
    table = ir.Table(rows=[headers, row1, row2], header_rows=1)

    sm = lm.SlideModel(id="tbl-test", archetype="table")
    _build_table_element(table, cr.x, cr.y, cr.w, cr.h, brand, sm)

    tbl_el = next(el for el in sm.elements if isinstance(el, lm.TableEl))
    assert len(tbl_el.rows) == 3

    # Header cells have subtle surface fill and distinct bottom divider
    for cell in tbl_el.rows[0]:
        assert cell.fill == brand.roles.surface
        assert cell.border_bottom == brand.roles.border
        assert cell.border_bottom_w >= 1.0
        assert getattr(cell, "is_header", False) is True

    # Alternating row fill
    assert tbl_el.rows[1][0].fill == brand.roles.surface


def test_table_render_generous_padding_and_borders():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    brand = palette_brand("ocean")

    headers = [lm.CellEl(paras=[lm.Para(runs=[lm.TRun(text="Metric", font="Calibri", size=14, color=brand.roles.text)])], is_header=True, border_bottom=brand.roles.border, border_bottom_w=1.5)]
    data = [lm.CellEl(paras=[lm.Para(runs=[lm.TRun(text="99.9%", font="Calibri", size=13, color=brand.roles.text)])], is_header=False, border_bottom=brand.roles.border, border_bottom_w=0.75)]
    tbl_el = lm.TableEl(
        name="Test Table",
        x=500000,
        y=500000,
        w=4000000,
        h=2000000,
        col_widths=[4000000],
        row_heights=[1000000, 1000000],
        rows=[headers, data],
    )
    render_table_element(slide, tbl_el)

    tbl_shape = slide.shapes[0]
    table = tbl_shape.table
    c0 = table.cell(0, 0)
    # Generous padding on header
    assert c0.margin_top >= 80000
    assert c0.margin_left >= 100000


def test_slide_kicker_accent_dot_and_letterspacing():
    brand = palette_brand("ocean")
    brand = apply_preset(brand, "corporate")

    slide = ir.Slide(
        index=1,
        title="Zero Trust Architecture",
        kicker="Enterprise Security Architecture",
        blocks=[ir.Paragraph(runs=[ir.Run(t="Body description")])],
    )
    sm = layout_slide(slide, brand, slide_num=1)

    title_el = next(el for el in sm.elements if isinstance(el, lm.TextEl) and el.name == "Slide Title")
    assert len(title_el.paras) >= 2

    kicker_para = title_el.paras[0]
    kicker_run = kicker_para.runs[0]
    assert kicker_run.text.startswith("• ")
    assert "ENTERPRISE SECURITY ARCHITECTURE" in kicker_run.text
    assert kicker_run.color == brand.roles.primary
    assert kicker_run.bold is True


def test_card_kicker_accent_dot():
    from htmlppt.layout.engine import _build_card_element

    brand = palette_brand("ocean")
    brand = apply_preset(brand, "corporate")
    sm = lm.SlideModel(id="card-test", archetype="cards")

    card = ir.Card(
        title="Sovereign Hosting",
        kicker="Pillar 1 · Sovereign Hosting",
        body=[ir.Run(t="Card body text")],
    )
    _build_card_element(card, 0, 0, 3000000, 3000000, "cards", brand, sm)

    content_el = next(el for el in sm.elements if isinstance(el, lm.TextEl) and el.name.startswith("Card Content"))
    kicker_para = content_el.paras[0]
    kicker_run = kicker_para.runs[0]
    assert kicker_run.text.startswith("• ")
    assert "PILLAR 1" in kicker_run.text
    assert kicker_run.color == brand.roles.primary
    assert kicker_run.bold is True
