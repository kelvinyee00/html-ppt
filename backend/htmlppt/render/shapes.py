"""Native PowerPoint vector shapes, fills and borders."""
from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.slide import Slide
from pptx.util import Emu, Pt

from htmlppt.brand.model import hex_to_rgb
from htmlppt.layout import model as lm
from htmlppt.render.text import ALIGN_MAP, VALIGN_MAP

GEOM_MAP = {
    "rect": MSO_SHAPE.RECTANGLE,
    "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "line": MSO_SHAPE.RECTANGLE,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "chevron": MSO_SHAPE.CHEVRON,
    "homePlate": MSO_SHAPE.PENTAGON,
}


def render_shape_element(slide: Slide, el: lm.ShapeEl) -> None:
    """Render a native PowerPoint auto shape."""
    shape_type = GEOM_MAP.get(el.geom, MSO_SHAPE.RECTANGLE)
    shape = slide.shapes.add_shape(shape_type, Emu(el.x), Emu(el.y), Emu(el.w), Emu(el.h))

    # Corner radius adjustment for rounded rectangles
    if el.geom == "roundRect" and len(shape.adjustments) > 0:
        shape.adjustments[0] = min(max(el.radius, 0.0), 0.5)

    # Fill (solid or gradient fallback)
    fill_hex = el.fill or (el.gradient[0] if el.gradient else None) or (el.line if el.geom == "line" else None)
    if fill_hex:
        shape.fill.solid()
        r, g, b = hex_to_rgb(fill_hex)
        shape.fill.fore_color.rgb = RGBColor(r, g, b)
    else:
        shape.fill.background()

    # Border line
    if el.geom != "line" and el.line and el.line_w > 0:
        shape.line.color.rgb = RGBColor(*hex_to_rgb(el.line))
        shape.line.width = Pt(el.line_w)
        if el.dash:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        shape.line.fill.background()

    # Shape text if present
    if el.paras:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = VALIGN_MAP.get(el.valign, MSO_ANCHOR.MIDDLE)
        tf.margin_left = Emu(el.inset)
        tf.margin_top = Emu(el.inset)
        tf.margin_right = Emu(el.inset)
        tf.margin_bottom = Emu(el.inset)

        for i, p_model in enumerate(el.paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN_MAP.get(p_model.align, PP_ALIGN.LEFT)
            if p_model.space_before > 0:
                p.space_before = Pt(p_model.space_before)
            if p_model.space_after > 0:
                p.space_after = Pt(p_model.space_after)
            for r_model in p_model.runs:
                run = p.add_run()
                run.text = r_model.text
                run.font.name = r_model.font
                run.font.size = Pt(r_model.size)
                run.font.bold = r_model.bold
                run.font.italic = r_model.italic
                run.font.underline = r_model.underline
                if r_model.color:
                    run.font.color.rgb = RGBColor(*hex_to_rgb(r_model.color))
