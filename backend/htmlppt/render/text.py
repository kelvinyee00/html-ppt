"""Native PowerPoint text element rendering."""
from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.slide import Slide
from pptx.util import Emu, Pt

from htmlppt.brand.model import hex_to_rgb
from htmlppt.layout import model as lm


ALIGN_MAP = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}

VALIGN_MAP = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}


def render_text_element(slide: Slide, el: lm.TextEl) -> None:
    """Create a native PowerPoint text box and render all paragraphs and runs."""
    tb = slide.shapes.add_textbox(Emu(el.x), Emu(el.y), Emu(el.w), Emu(el.h))
    tf = tb.text_frame
    tf.word_wrap = el.wrap
    tf.vertical_anchor = VALIGN_MAP.get(el.valign, MSO_ANCHOR.TOP)
    tf.margin_left = Emu(el.inset)
    tf.margin_top = Emu(el.inset)
    tf.margin_right = Emu(el.inset)
    tf.margin_bottom = Emu(el.inset)

    # Optional background fill for text box
    if el.fill:
        tb.fill.solid()
        r, g, b = hex_to_rgb(el.fill)
        tb.fill.fore_color.rgb = RGBColor(r, g, b)

    for i, p_model in enumerate(el.paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN_MAP.get(p_model.align, PP_ALIGN.LEFT)
        p.level = p_model.level

        if p_model.space_before > 0:
            p.space_before = Pt(p_model.space_before)
        if p_model.space_after > 0:
            p.space_after = Pt(p_model.space_after)

        # Bullet character styling
        _apply_bullet(p, p_model)

        for r_model in p_model.runs:
            run = p.add_run()
            run.text = r_model.text
            run.font.name = r_model.font
            run.font.size = Pt(r_model.size)
            run.font.bold = r_model.bold
            run.font.italic = r_model.italic
            run.font.underline = r_model.underline

            if r_model.color:
                r, g, b = hex_to_rgb(r_model.color)
                run.font.color.rgb = RGBColor(r, g, b)

            if r_model.href:
                run.hyperlink.address = r_model.href


_A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
_BULLET_INDENT_EMU = 228600  # 0.25in hanging indent per level


def _apply_bullet(p, p_model: lm.Para) -> None:
    """Write bullet XML for a paragraph: none, auto-numbered ("1.") or a bullet character."""
    pPr = p._p.get_or_add_pPr()
    if p_model.bullet is None:
        pPr.append(parse_xml(f"<a:buNone {_A_NS}/>"))
        return

    pPr.set("marL", str(_BULLET_INDENT_EMU * (p_model.level + 1)))
    pPr.set("indent", str(-_BULLET_INDENT_EMU))
    if p_model.bullet_color:
        clean_bc = p_model.bullet_color.lstrip("#")
        pPr.append(parse_xml(f'<a:buClr {_A_NS}><a:srgbClr val="{clean_bc}"/></a:buClr>'))
    if p_model.bullet.rstrip(".").isdigit():
        pPr.append(parse_xml(f'<a:buAutoNum {_A_NS} type="arabicPeriod"/>'))
    else:
        char = p_model.bullet.replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;")
        pPr.append(parse_xml(f'<a:buChar {_A_NS} char="{char}"/>'))

