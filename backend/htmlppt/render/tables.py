"""Native PowerPoint table element rendering."""
from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.slide import Slide
from pptx.util import Emu, Pt

from htmlppt.brand.model import hex_to_rgb
from htmlppt.layout import model as lm
from htmlppt.render.text import ALIGN_MAP


def render_table_element(slide: Slide, el: lm.TableEl) -> None:
    """Render a native, editable PowerPoint table with clean horizontal divider rules and padded cells."""
    n_rows = len(el.rows)
    n_cols = len(el.col_widths)
    if not n_rows or not n_cols:
        return

    table_shape = slide.shapes.add_table(n_rows, n_cols, Emu(el.x), Emu(el.y), Emu(el.w), Emu(el.h))
    table = table_shape.table

    # Set column widths
    for ci, cw in enumerate(el.col_widths):
        table.columns[ci].width = Emu(cw)

    # Set row heights
    for ri, rh in enumerate(el.row_heights):
        table.rows[ri].height = Emu(rh)

    # Populate cells
    for ri, row_data in enumerate(el.rows):
        for ci, cell_data in enumerate(row_data):
            if ci >= n_cols:
                continue
            cell = table.cell(ri, ci)
            tc_pr = cell._tc.get_or_add_tcPr()

            # Background fill
            if cell_data.fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(*hex_to_rgb(cell_data.fill))

            # Generous cell padding: header rows get extra breathing room
            is_head = getattr(cell_data, "is_header", False) or (ri == 0)
            m_horiz = Emu(120000 if is_head else 100000)
            m_vert = Emu(90000 if is_head else 70000)
            cell.margin_left = m_horiz
            cell.margin_right = m_horiz
            cell.margin_top = m_vert
            cell.margin_bottom = m_vert

            tf = cell.text_frame
            tf.word_wrap = True
            tf.margin_left = m_horiz
            tf.margin_right = m_horiz
            tf.margin_top = m_vert
            tf.margin_bottom = m_vert

            # Modern horizontal divider rules - eliminate vertical grid lines
            _apply_cell_borders(tc_pr, cell_data)

            for pi, p_model in enumerate(cell_data.paras):
                p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                p.alignment = ALIGN_MAP.get(p_model.align, PP_ALIGN.LEFT)

                for r_model in p_model.runs:
                    run = p.add_run()
                    run.text = r_model.text
                    run.font.name = r_model.font
                    run.font.size = Pt(r_model.size)
                    run.font.bold = r_model.bold
                    run.font.italic = r_model.italic
                    if r_model.color:
                        run.font.color.rgb = RGBColor(*hex_to_rgb(r_model.color))


def _apply_cell_borders(tc_pr, cell_data: lm.CellEl) -> None:
    """Apply horizontal divider rules and remove harsh vertical gridlines."""
    # Clear and disable vertical gridlines
    for edge in ("lnL", "lnR"):
        for child in list(tc_pr):
            if child.tag.endswith(edge):
                tc_pr.remove(child)
        xml = f'<a:{edge} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:noFill/></a:{edge}>'
        tc_pr.append(parse_xml(xml))

    # Clean horizontal divider rule on bottom
    for child in list(tc_pr):
        if child.tag.endswith("lnB"):
            tc_pr.remove(child)

    if cell_data.border_bottom:
        clean_color = cell_data.border_bottom.lstrip("#")
        w_emu = int((cell_data.border_bottom_w or 1.0) * 12700)
        xml = f'<a:lnB xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" w="{w_emu}" cmpd="s"><a:solidFill><a:srgbClr val="{clean_color}"/></a:solidFill></a:lnB>'
        tc_pr.append(parse_xml(xml))
    else:
        xml = '<a:lnB xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:noFill/></a:lnB>'
        tc_pr.append(parse_xml(xml))

